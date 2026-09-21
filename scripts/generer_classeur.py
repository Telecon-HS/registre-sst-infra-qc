"""Génère le classeur du registre (15 onglets) à partir de donnees/classeur/.

    python scripts/generer_classeur.py        → sortie/Registre_risques_SST_Falcon_Infra_QC.xlsx

Chaque onglet est décrit par un fichier JSON : cellules (valeur ou formule, style),
largeurs, hauteurs, fusions, volets figés, filtres et listes déroulantes.
Les styles sont partagés dans _styles.json. Les formules sont conservées telles
quelles : ouvrir le fichier dans Excel ou LibreOffice les recalcule.

Pour modifier le contenu : éditer la valeur de la cellule dans le JSON de l'onglet.
"""
import datetime as dt
from copy import copy

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from commun import (DONNEES, PRIVE, RACINE, SORTIE, ligne_version, lire_json,
                    nom_versionne, prive_disponible, remplacer_jetons, table_des_noms)

CLASSEUR = DONNEES / "classeur"


def couleur(c):
    if not c:
        return None
    if "rgb" in c:
        return Color(rgb=c["rgb"])
    if "theme" in c:
        return Color(theme=c["theme"], tint=c["tint"])
    return Color(indexed=c["indexed"])


def construire_styles():
    styles = []
    for s in lire_json(CLASSEUR / "_styles.json"):
        f = s["font"]
        font = Font(name=f["name"], sz=f["sz"], b=f["b"], i=f["i"], u=f["u"], strike=f["strike"], color=couleur(f["color"]))
        fill = PatternFill(fill_type=s["fill"]["type"], fgColor=couleur(s["fill"]["fg"]), bgColor=couleur(s["fill"]["bg"])) if s["fill"] else PatternFill()
        cote = lambda d: Side(style=d["style"], color=couleur(d["color"])) if d else Side()  # noqa: E731
        border = Border(**{k: cote(s["border"][k]) for k in ("left", "right", "top", "bottom")})
        a = s["align"]
        align = Alignment(horizontal=a["h"], vertical=a["v"], wrap_text=a["wrap"], indent=a["indent"], text_rotation=a["rot"], shrink_to_fit=a["shrink"])
        styles.append((font, fill, border, align, s["numfmt"]))
    return styles


def valeur(v):
    if isinstance(v, dict) and "date" in v:
        return dt.datetime.fromisoformat(v["date"])
    return v


def onglet_volumes(wb, styles):
    """Onglet « Volumes et tendances » : les trois relevés du même volume,
    avec une courbe par type de formulaire et une comparaison des sources."""
    V = lire_json(DONNEES / "volumes.json")
    ws = wb.create_sheet("Volumes et tendances")
    ws.sheet_view.showGridLines = False
    titre, gras, normal, petit = styles

    ws["B2"] = "Volumes d’activité — trois relevés, trois portées"
    ws["B2"]._style = copy(titre._style)
    ws["B3"] = V["_note"]
    ws["B3"]._style = copy(petit._style)
    for i, (k, t) in enumerate(V["sources"].items(), start=4):
        ws[f"B{i}"] = t
        ws[f"B{i}"]._style = copy(petit._style)

    # --- tableau mensuel (export daté)
    d = 9
    ws[f"B{d}"] = "Volumes mensuels — export daté de 165 fiches"
    ws[f"B{d}"]._style = copy(gras._style)
    ws.cell(d + 1, 2, "Type de formulaire")._style = copy(gras._style)
    for j, m in enumerate(V["mois"]):
        c = ws.cell(d + 1, 3 + j, m[5:] + "/" + m[2:4])
        c._style = copy(gras._style)
    ws.cell(d + 1, 3 + len(V["mois"]), "Total")._style = copy(gras._style)
    for i, l in enumerate(V["lignes"]):
        r = d + 2 + i
        ws.cell(r, 2, l["type"])._style = copy(normal._style)
        for j, n in enumerate(l["mois"]):
            ws.cell(r, 3 + j, n)._style = copy(normal._style)
        ws.cell(r, 3 + len(V["mois"]), f"=SUM(C{r}:{chr(66 + len(V['mois']))}{r})")._style = copy(gras._style)
    fin = d + 1 + len(V["lignes"])

    courbe = LineChart()
    courbe.title = "Tendance mensuelle par type de formulaire — export daté"
    courbe.y_axis.title = "Fiches"
    courbe.height, courbe.width = 9, 26
    courbe.add_data(Reference(ws, min_col=2, max_col=2 + len(V["mois"]), min_row=d + 2, max_row=fin), from_rows=True, titles_from_data=True)
    courbe.set_categories(Reference(ws, min_col=3, max_col=2 + len(V["mois"]), min_row=d + 1))
    for serie in courbe.series:
        serie.smooth = False
    ws.add_chart(courbe, f"B{fin + 2}")

    # --- comparaison des trois relevés
    c0 = fin + 22
    ws[f"B{c0}"] = "Le même volume, vu par trois relevés"
    ws[f"B{c0}"]._style = copy(gras._style)
    entetes = ["Type d’activité", "Registre v7 (19 mai – 17 sept.)", "Export daté (2026, 165 fiches)", "Tableau de bord (2026)"]
    for j, t in enumerate(entetes):
        ws.cell(c0 + 1, 2 + j, t)._style = copy(gras._style)
    for i, l in enumerate(V["comparaison"]):
        r = c0 + 2 + i
        ws.cell(r, 2, l["type"])._style = copy(normal._style)
        for j, k in enumerate(["registre", "export", "tableau"]):
            ws.cell(r, 3 + j, l[k])._style = copy(normal._style)

    barres = BarChart()
    barres.type, barres.title = "col", "Écart entre les trois relevés"
    barres.height, barres.width = 9, 26
    barres.add_data(Reference(ws, min_col=3, max_col=5, min_row=c0 + 1, max_row=c0 + 1 + len(V["comparaison"])), titles_from_data=True)
    barres.set_categories(Reference(ws, min_col=2, min_row=c0 + 2, max_row=c0 + 1 + len(V["comparaison"])))
    ws.add_chart(barres, f"B{c0 + 3 + len(V['comparaison'])}")

    bas = c0 + 23
    ws[f"B{bas}"] = "Lecture — " + V["lecture"]
    ws[f"B{bas}"]._style = copy(normal._style)
    ws[f"B{bas + 2}"] = "Réserve — " + V["reserve"]
    ws[f"B{bas + 2}"]._style = copy(petit._style)
    for lettre, largeur in [("A", 2.5), ("B", 52)]:
        ws.column_dimensions[lettre].width = largeur
    for j in range(len(V["mois"])):
        ws.column_dimensions[chr(67 + j)].width = 9


def generer(avec_prive=True, sortie=None):
    avec_prive = avec_prive and prive_disponible()
    table = table_des_noms(avec_prive)
    cellules_privees = lire_json(PRIVE / "classeur_prive.json") if avec_prive and (PRIVE / "classeur_prive.json").exists() else {}
    styles = construire_styles()

    wb = Workbook()
    wb.remove(wb.active)
    for nom in lire_json(CLASSEUR / "_ordre.json"):
        f = lire_json(CLASSEUR / nom)
        ws = wb.create_sheet(f["titre"])
        for coord, v, si in f["cellules"]:
            c = ws[coord]
            cle = f"{f['titre']}!{coord}"
            v = cellules_privees.get(cle, v)
            if isinstance(v, str) and "__DATE_VERSION__" in v:      # version unique : donnees/version.json
                ver = lire_json(DONNEES / "version.json")
                v = v.replace("__DATE_VERSION__", f"{ver['date_longue']} — version {ver['version']}")
            c.value = valeur(remplacer_jetons(v, table))
            c.font, c.fill, c.border, c.alignment, c.number_format = styles[si]
        for col, largeur in f["colonnes"].items():
            ws.column_dimensions[col].width = largeur
        for lig, hauteur in f["lignes"].items():
            ws.row_dimensions[int(lig)].height = hauteur
        for plage in f["fusions"]:
            ws.merge_cells(plage)
        ws.sheet_view.showGridLines = f["grille"]
        if f["figer"]:
            ws.freeze_panes = f["figer"]
        if f["filtre"]:
            ws.auto_filter.ref = f["filtre"]
        for v in f["validations"]:
            dv = DataValidation(type=v["type"], formula1=v["formula1"], allow_blank=v["allow_blank"])
            dv.add(v["sqref"])
            ws.add_data_validation(dv)
        for im in f["images"]:
            img = Image(str(RACINE / im["fichier"]))
            img.width, img.height = im["largeur"], im["hauteur"]
            ws.add_image(img, im["ancre"])

    g = wb["Garde"]
    onglet_volumes(wb, (g["B2"], g["B20"] if g["B20"].value else g["B2"], g["C20"] if g["C20"].value else g["B34"], g["B34"]))

    garde = wb["Garde"]
    modele = garde["B34"]            # note de bas de page de la garde
    garde.merge_cells("B38:C39")
    garde["B38"].value = ligne_version("classeur — source de travail")
    garde["B38"]._style = copy(modele._style)

    sortie = sortie or SORTIE / nom_versionne("Registre_risques_SST_Falcon_Infra_QC", ".xlsx")
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)
    print(f"classeur → {sortie} [{'complet' if avec_prive else 'sans données privées'}]")
    return sortie


if __name__ == "__main__":
    import sys
    generer(avec_prive="--sans-prive" not in sys.argv)
