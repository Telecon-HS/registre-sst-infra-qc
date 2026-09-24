"""Tests de l'onglet « Dossiers transférés au comité ». Lancer : pytest -q"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import generer_classeur as gc  # noqa: E402
from commun import DONNEES, PRIVE  # noqa: E402

D = json.loads((DONNEES / "dossiers_comite.json").read_text(encoding="utf-8"))


def test_le_dossier_verse_est_decrit():
    d = next(x for x in D["dossiers"] if x["numero"] == "29384431")
    assert d["date_evenement"] == "2026-08-19"
    assert d["etat"] in D["etats"]
    assert d["etat"].startswith("ouvert")          # l'enquête n'est pas terminée


def test_compte_des_pieces():
    l = gc.lignes_dossiers(D)[0]
    assert l["pieces"] == "3 sur 7" and l["a_obtenir"] == 4


def test_valeurs_inconnues_marquees_a_confirmer():
    l = gc.lignes_dossiers(D)[0]
    assert l["accuse"] == "à confirmer" and l["date_accuse"] == "à confirmer"
    assert l["risque"] == "à confirmer"


def test_aucun_nom_ni_donnee_medicale():
    texte = json.dumps(D, ensure_ascii=False)
    f = PRIVE / "noms.json"
    if f.exists():
        noms = [g for gs in json.loads(f.read_text(encoding="utf-8"))["variantes"].values() for g in gs]
        assert not [n for n in noms if n in texte]
    for mot in ("diagnostic médical", "physiothérapie", "chiropraxie", "médecin traitant"):
        assert mot not in texte
    assert "⟦P01⟧" in texte                         # les personnes passent par des jetons


def test_ce_qui_ne_se_transfere_pas_est_liste_sans_son_contenu():
    d = D["dossiers"][0]
    elements = [x["element"] for x in d["ne_se_transfere_pas"]]
    assert "Pièces médicales, diagnostic, suivi thérapeutique" in elements
    assert all("motif" in x and "responsable" in x for x in d["ne_se_transfere_pas"])


# --- Mesures correctives du dossier 29384431 ---
import datetime as dt  # noqa: E402

M = next(x for x in D["dossiers"] if x["numero"] == "29384431")["mesures"]


def test_huit_mesures_completes():
    assert [m["numero"] for m in M] == list(range(1, 9))
    for m in M:
        assert m["type_controle"] in D["types_controle"]
        assert m["responsable"].startswith("⟦P") or m["responsable"] == "à désigner"
        dt.date.fromisoformat(m["date_visee"])


def test_etat_jamais_deduit_d_une_date():
    for m in M:
        assert m["etat"] in D["etats_mesure"]
        if m["etat_inscrit_par"] is None:
            assert m["etat"] == "à confirmer"
    # même avec toutes les dates visées passées, l'état ne bouge pas
    lignes = gc.lignes_mesures(D)
    assert all(l["etat"] == "à confirmer" for l in lignes)


def test_mesures_4_et_5_portent_le_plan_et_attendent_un_responsable():
    for m in M:
        assert bool(m.get("porte_le_plan")) == (m["numero"] in (4, 5))
    for m in (M[3], M[4]):
        assert m["responsable"] == "à désigner" and m["prealable"]


def test_vue_en_retard():
    lignes = gc.lignes_mesures(D)
    v = gc.vues_suivi(lignes, dt.date(2026, 9, 24))
    assert [l["numero"] for l in v["en_retard"]] == [1, 6, 7, 8]
    v = gc.vues_suivi(lignes, dt.date(2026, 9, 15))
    assert v["en_retard"] == []                       # la date visée elle-même n'est pas un retard
    lignes[0]["etat"] = "réalisée"
    v = gc.vues_suivi(lignes, dt.date(2026, 9, 24))
    assert 1 not in [l["numero"] for l in v["en_retard"]]


def test_vue_par_responsable_ne_confond_pas_les_replis():
    v = gc.vues_suivi(gc.lignes_mesures(D), dt.date(2026, 9, 24))
    assert {k: [l["numero"] for l in g] for k, g in v["par_responsable"].items()} == {
        "⟦P06⟧": [1, 6], "⟦P18⟧": [2, 8], "⟦P21⟧": [3, 7], "à désigner": [4, 5]}


def test_vue_efficacite_non_verifiee():
    lignes = gc.lignes_mesures(D)
    assert gc.vues_suivi(lignes, dt.date(2026, 9, 24))["non_verifiee"] == []
    lignes[1]["etat"] = "réalisée"
    lignes[2]["etat"], lignes[2]["efficacite"] = "réalisée", "vérifiée le 2026-11-01"
    assert [l["numero"] for l in gc.vues_suivi(lignes, dt.date(2026, 9, 24))["non_verifiee"]] == [2]


def test_onglet_de_suivi_genere(tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(gc.generer(avec_prive=False, sortie=tmp_path / "c.xlsx"))
    ws = wb["Suivi des mesures correctives"]
    texte = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "En retard" in texte and "Par responsable" in texte and "Efficacité non vérifiée" in texte
    assert "⟦" not in texte


# --- Mode séance de la page Comité ---
import generer_page as gp  # noqa: E402
from commun import GABARITS  # noqa: E402


def test_page_recoit_le_dossier_sans_retard_precalcule():
    d = next(x for x in gp.dossiers_pour_la_page() if x["numero"] == "29384431")
    assert len(d["mesures"]) == 8 and len(d["points"]) == 5 and d["accuse_en_attente"]
    assert set(d["mesures"][0]) == {"numero", "date_visee", "etat"}   # le retard se calcule dans la page, au jour de la séance
    assert "en_retard" not in d


def test_gabarit_mode_seance_bloc_dossiers():
    g = (GABARITS / "page_registre.html").read_text(encoding="utf-8")
    debut = g.index("const blocDossiers")
    bloc = g[debut:g.index("document.getElementById('seance').innerHTML", debut)]
    assert "Ce que le comité doit décider" in bloc and "en retard" in bloc
    assert "m.etat !== 'réalisée'" in g                      # même règle que le classeur et les alertes
    assert "#" not in bloc.replace("${", "")                 # aucune couleur codée en dur : variables du thème seulement
    assert "#seance .s-gauche{display:block;overflow:visible}" in g   # affichage téléphone


def test_page_terrain_ne_recoit_pas_les_dossiers(tmp_path):
    t = gp.generer_terrain(tmp_path / "terrain.html").read_text(encoding="utf-8")
    assert "29384431" not in t and "dossiers_transferes" not in t
